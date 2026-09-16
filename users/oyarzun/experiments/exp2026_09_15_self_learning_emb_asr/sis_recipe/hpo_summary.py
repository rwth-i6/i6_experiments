import os
import re
import json
import statistics
from typing import Dict, List, Any, Optional

from sisyphus import tk, Task
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


class HpoResultsExcelJob(tk.Job):
    """
    Sisyphus Job that aggregates PER evaluations, variance metrics, codebook perplexities,
    and runtimes across all HPO training configurations and compiles:
    1. An Excel report and CSV summary with explicit Killed/Pending/Completed status.
    2. Automated perplexity trajectory plots with mean curves and confidence bounds.
    3. Automated representation variance plots (Audio, Text, Mixed) with confidence bounds.
    """

    def __init__(
        self,
        *,
        configs_meta: Dict[str, Dict[str, Any]],
        train_jobs: Dict[str, Any],
        eval_score_results: Optional[Dict[str, Dict[int, tk.Path]]] = None,
        variance_outputs: Optional[Dict[str, Dict[int, tk.Path]]] = None,
        eval_epochs: Optional[List[int]] = None,
        target_num_epochs: int = 200,
        prefix_name: Optional[str] = None,
    ):
        super().__init__()
        self.configs_meta = configs_meta
        self.train_jobs = train_jobs
        self.eval_score_results = eval_score_results or {}
        self.variance_outputs = variance_outputs or {}
        self.eval_epochs = eval_epochs or [50, 100, 150, 200]
        self.target_num_epochs = target_num_epochs

        self.out_excel = self.output_path("v6_codebook_opt_summary.xlsx")
        self.out_csv = self.output_path("v6_codebook_opt_summary.csv")
        self.out_ppl_plot = self.output_path("v6_codebook_opt_perplexity_summary.png")
        self.out_var_plot = self.output_path("v6_codebook_opt_variance_summary.png")

        if prefix_name:
            self.add_alias(f"{prefix_name}/joined_results/hpo_excel_job")
            tk.register_output(f"{prefix_name}/joined_results/v6_codebook_opt_summary.xlsx", self.out_excel)
            tk.register_output(f"{prefix_name}/joined_results/v6_codebook_opt_summary.csv", self.out_csv)
            tk.register_output(f"{prefix_name}/joined_results/v6_codebook_opt_perplexity_summary.png", self.out_ppl_plot)
            tk.register_output(f"{prefix_name}/joined_results/v6_codebook_opt_variance_summary.png", self.out_var_plot)

    def tasks(self):
        yield Task("run", mini_task=True)

    def run(self):
        time_pattern = re.compile(r"':meta:epoch_train_time_secs':\s*([0-9.]+)")

        records = []
        all_prob_ppl_curves = {}
        all_hard_ppl_curves = {}

        # For tracking variance across epochs: ep -> list of values across all configs
        epoch_audio_vars = {ep: [] for ep in self.eval_epochs}
        epoch_text_vars = {ep: [] for ep in self.eval_epochs}
        epoch_mixed_vars = {ep: [] for ep in self.eval_epochs}

        for name, meta in self.configs_meta.items():
            runtimes = []
            prob_ppls = {}
            hard_ppls = {}
            max_epoch_seen = 0
            is_killed = False
            is_completed = False

            # 1. Parse learning_rates & determine status
            if name in self.train_jobs:
                t_job = self.train_jobs[name]
                lr_path = None
                train_job_dir = None
                try:
                    if hasattr(t_job, "out_files") and "learning_rates" in t_job.out_files:
                        lr_path = t_job.out_files["learning_rates"].get_path()
                        train_job_dir = os.path.dirname(os.path.dirname(os.path.realpath(lr_path)))
                    elif hasattr(t_job, "output_path"):
                        train_job_dir = t_job.output_path().get_path()
                        cand = os.path.join(train_job_dir, "learning_rates")
                        if os.path.exists(cand):
                            lr_path = cand
                except Exception:
                    pass

                if lr_path and os.path.exists(lr_path):
                    try:
                        with open(lr_path, "r") as f:
                            content = f.read()

                        for line in content.splitlines():
                            m_time = time_pattern.search(line)
                            if m_time:
                                runtimes.append(float(m_time.group(1)))

                        epoch_blocks = re.split(r"(?m)^(?=\d+:\s*EpochData)", content)
                        for block in epoch_blocks:
                            m_ep = re.match(r"^(\d+):", block.strip())
                            if not m_ep:
                                continue
                            ep = int(m_ep.group(1))
                            if ep > max_epoch_seen:
                                max_epoch_seen = ep

                            m_prob = re.search(r"'(?:train_)?(?:error_)?codebook_prob_ppl[^']*':\s*([0-9.]+)", block)
                            if m_prob:
                                prob_ppls[ep] = float(m_prob.group(1))

                            m_hard = re.search(r"'(?:train_)?(?:error_)?codebook_hard_ppl[^']*':\s*([0-9.]+)", block)
                            if m_hard:
                                hard_ppls[ep] = float(m_hard.group(1))
                    except Exception:
                        pass

                # Check if killed / finished
                if max_epoch_seen >= self.target_num_epochs:
                    is_completed = True
                elif train_job_dir and os.path.isdir(train_job_dir):
                    for f in os.listdir(train_job_dir):
                        if f.startswith(".killed") or f.startswith(".interrupted") or f.startswith(".error"):
                            is_killed = True
                            break

                if not is_completed and max_epoch_seen > 0 and (is_killed or max_epoch_seen < self.target_num_epochs):
                    status_str = f"Killed (Ep {max_epoch_seen})" if is_killed else f"Running (Ep {max_epoch_seen})"
                elif is_completed:
                    status_str = "Completed"
                else:
                    status_str = "Pending"
            else:
                status_str = "Pending"

            if prob_ppls:
                all_prob_ppl_curves[name] = prob_ppls
            if hard_ppls:
                all_hard_ppl_curves[name] = hard_ppls

            # 2. Parse PER
            pers = {}
            for ep in self.eval_epochs:
                if ep > max_epoch_seen:
                    pers[ep] = "Killed" if is_killed else "Pending"
                else:
                    pers[ep] = ""

            if name in self.eval_score_results:
                for ep, path_obj in self.eval_score_results[name].items():
                    try:
                        p = path_obj.get_path() if hasattr(path_obj, "get_path") else str(path_obj)
                        if os.path.isfile(p):
                            with open(p, "r") as f:
                                pers[ep] = f.read().strip()
                    except Exception:
                        pass

            # 3. Parse Variance Metrics
            audio_vars = {ep: ("Killed" if (is_killed and ep > max_epoch_seen) else ("Pending" if ep > max_epoch_seen else "")) for ep in self.eval_epochs}
            text_vars = {ep: ("Killed" if (is_killed and ep > max_epoch_seen) else ("Pending" if ep > max_epoch_seen else "")) for ep in self.eval_epochs}
            mixed_vars = {ep: ("Killed" if (is_killed and ep > max_epoch_seen) else ("Pending" if ep > max_epoch_seen else "")) for ep in self.eval_epochs}

            if name in self.variance_outputs:
                for ep, path_obj in self.variance_outputs[name].items():
                    try:
                        base_p = path_obj.get_path() if hasattr(path_obj, "get_path") else str(path_obj)
                        summary_file = os.path.join(base_p, "variance_summary.txt")
                        if not os.path.exists(summary_file) and os.path.isfile(base_p):
                            summary_file = base_p

                        if os.path.isfile(summary_file):
                            with open(summary_file, "r") as f:
                                c = f.read()
                            m_a = re.search(r"Audio variance.*:\s*([0-9.]+)", c)
                            m_t = re.search(r"Text variance.*:\s*([0-9.]+)", c)
                            m_m = re.search(r"Mixed variance.*:\s*([0-9.]+)", c)
                            if m_a:
                                val_a = float(m_a.group(1))
                                audio_vars[ep] = f"{val_a:.4f}"
                                epoch_audio_vars[ep].append(val_a)
                            if m_t:
                                val_t = float(m_t.group(1))
                                text_vars[ep] = f"{val_t:.4f}"
                                epoch_text_vars[ep].append(val_t)
                            if m_m:
                                val_m = float(m_m.group(1))
                                mixed_vars[ep] = f"{val_m:.4f}"
                                epoch_mixed_vars[ep].append(val_m)
                    except Exception:
                        pass

            mean_rt = round(statistics.mean(runtimes), 2) if runtimes else ""
            total_rt = round(sum(runtimes) / 3600.0, 2) if runtimes else ""
            latest_prob_ppl = round(prob_ppls[max(prob_ppls.keys())], 2) if prob_ppls else ("Killed" if is_killed else "")
            latest_hard_ppl = round(hard_ppls[max(hard_ppls.keys())], 2) if hard_ppls else ("Killed" if is_killed else "")

            records.append({
                "name": name,
                "status": status_str,
                "layers": meta.get("layers", ""),
                "cb_prob": meta.get("cb_prob", ""),
                "cb_vars": meta.get("cb_vars", ""),
                "cb_groups": meta.get("cb_groups", ""),
                "div_loss": meta.get("div_loss", ""),
                "mean_rt": mean_rt,
                "total_rt": total_rt,
                "latest_prob_ppl": latest_prob_ppl,
                "latest_hard_ppl": latest_hard_ppl,
                "prob_ppls": prob_ppls,
                "hard_ppls": hard_ppls,
                "pers": pers,
                "audio_vars": audio_vars,
                "text_vars": text_vars,
                "mixed_vars": mixed_vars,
                "is_killed": is_killed,
            })

        # --- CSV Generation ---
        headers = [
            "Config Name", "Status", "Layers", "CB Prob", "CB Vars", "CB Groups", "Div Scale",
            "Mean Ep Time (s)", "Total Time (h)", "Latest Prob PPL", "Latest Hard PPL",
        ]
        for ep in self.eval_epochs:
            headers.append(f"Prob PPL Ep {ep}")
        for ep in self.eval_epochs:
            headers.append(f"Hard PPL Ep {ep}")
        for ep in self.eval_epochs:
            headers.append(f"PER Ep {ep}")
        for ep in self.eval_epochs:
            headers.append(f"Audio Var Ep {ep}")
        for ep in self.eval_epochs:
            headers.append(f"Text Var Ep {ep}")
        for ep in self.eval_epochs:
            headers.append(f"Mixed Var Ep {ep}")

        with open(self.out_csv.get_path(), "w") as f:
            f.write(",".join(headers) + "\n")
            for r in records:
                row = [
                    str(r["name"]), str(r["status"]), str(r["layers"]), str(r["cb_prob"]),
                    str(r["cb_vars"]), str(r["cb_groups"]), str(r["div_loss"]), str(r["mean_rt"]),
                    str(r["total_rt"]), str(r["latest_prob_ppl"]), str(r["latest_hard_ppl"]),
                ]
                for ep in self.eval_epochs:
                    val = str(r["prob_ppls"].get(ep, "Killed" if r["is_killed"] else ""))
                    row.append(val)
                for ep in self.eval_epochs:
                    val = str(r["hard_ppls"].get(ep, "Killed" if r["is_killed"] else ""))
                    row.append(val)
                for ep in self.eval_epochs:
                    row.append(str(r["pers"].get(ep, "")))
                for ep in self.eval_epochs:
                    row.append(str(r["audio_vars"].get(ep, "")))
                for ep in self.eval_epochs:
                    row.append(str(r["text_vars"].get(ep, "")))
                for ep in self.eval_epochs:
                    row.append(str(r["mixed_vars"].get(ep, "")))
                f.write(",".join(row) + "\n")

        # --- Styled Excel Generation ---
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "v6_Codebook_Opt_Summary"
        ws.append(headers)

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        killed_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        killed_font = Font(name="Calibri", size=10, color="C00000", italic=True)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, r in enumerate(records, start=2):
            row = [
                r["name"], r["status"], r["layers"], r["cb_prob"],
                r["cb_vars"], r["cb_groups"], r["div_loss"], r["mean_rt"],
                r["total_rt"], r["latest_prob_ppl"], r["latest_hard_ppl"],
            ]
            for ep in self.eval_epochs:
                row.append(r["prob_ppls"].get(ep, "Killed" if r["is_killed"] else ""))
            for ep in self.eval_epochs:
                row.append(r["hard_ppls"].get(ep, "Killed" if r["is_killed"] else ""))
            for ep in self.eval_epochs:
                row.append(r["pers"].get(ep, ""))
            for ep in self.eval_epochs:
                row.append(r["audio_vars"].get(ep, ""))
            for ep in self.eval_epochs:
                row.append(r["text_vars"].get(ep, ""))
            for ep in self.eval_epochs:
                row.append(r["mixed_vars"].get(ep, ""))
            ws.append(row)

            if r["is_killed"] or "Killed" in str(r["status"]):
                status_cell = ws.cell(row=row_idx, column=2)
                status_cell.fill = killed_fill
                status_cell.font = killed_font

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(self.out_excel.get_path())

        # --- 1. Automated Perplexity Summary Plot with Confidence Bounds ---
        try:
            plt.figure(figsize=(14, 6))

            # Subplot 1: Soft Probability Perplexity
            plt.subplot(1, 2, 1)
            # Find all recorded epochs
            all_eps = sorted(list({ep for c in all_prob_ppl_curves.values() for ep in c.keys()}))
            if all_eps:
                ep_means = [np.mean([c[e] for c in all_prob_ppl_curves.values() if e in c]) for e in all_eps]
                ep_stds = [np.std([c[e] for c in all_prob_ppl_curves.values() if e in c]) for e in all_eps]
                ep_means, ep_stds = np.array(ep_means), np.array(ep_stds)

                plt.plot(all_eps, ep_means, color="navy", linewidth=2.5, label="Mean Prob PPL")
                plt.fill_between(all_eps, ep_means - ep_stds, ep_means + ep_stds, color="blue", alpha=0.2, label="Confidence (±1 std)")

                for name, curve in all_prob_ppl_curves.items():
                    short_name = name.replace("asr_v6_codebook_opt_", "").replace("_ep-200", "")
                    plt.plot(sorted(curve.keys()), [curve[e] for e in sorted(curve.keys())], alpha=0.3, linewidth=0.8)

            plt.xlabel("Epoch")
            plt.ylabel("Codebook Soft Prob Perplexity")
            plt.title("Soft Codebook Perplexity Mean & Confidence Envelope")
            plt.grid(True, linestyle="--", alpha=0.6)
            plt.legend(loc="lower right")

            # Subplot 2: Hard Codebook Perplexity
            plt.subplot(1, 2, 2)
            all_hard_eps = sorted(list({ep for c in all_hard_ppl_curves.values() for ep in c.keys()}))
            if all_hard_eps:
                hard_means = [np.mean([c[e] for c in all_hard_ppl_curves.values() if e in c]) for e in all_hard_eps]
                hard_stds = [np.std([c[e] for c in all_hard_ppl_curves.values() if e in c]) for e in all_hard_eps]
                hard_means, hard_stds = np.array(hard_means), np.array(hard_stds)

                plt.plot(all_hard_eps, hard_means, color="darkgreen", linewidth=2.5, label="Mean Hard PPL")
                plt.fill_between(all_hard_eps, hard_means - hard_stds, hard_means + hard_stds, color="green", alpha=0.2, label="Confidence (±1 std)")

                for name, curve in all_hard_ppl_curves.items():
                    plt.plot(sorted(curve.keys()), [curve[e] for e in sorted(curve.keys())], alpha=0.3, linewidth=0.8)

            plt.xlabel("Epoch")
            plt.ylabel("Codebook Hard Perplexity")
            plt.title("Hard Codebook Perplexity Mean & Confidence Envelope")
            plt.grid(True, linestyle="--", alpha=0.6)
            plt.legend(loc="lower right")

            plt.tight_layout()
            plt.savefig(self.out_ppl_plot.get_path(), dpi=300, bbox_inches="tight")
            plt.close()
        except Exception:
            pass

        # --- 2. Automated Representation Variance Plot with Confidence Bounds ---
        try:
            plt.figure(figsize=(10, 6))

            eps = sorted([ep for ep in self.eval_epochs if epoch_audio_vars[ep] or epoch_text_vars[ep] or epoch_mixed_vars[ep]])
            if eps:
                a_means = [np.mean(epoch_audio_vars[e]) if epoch_audio_vars[e] else np.nan for e in eps]
                a_stds = [np.std(epoch_audio_vars[e]) if len(epoch_audio_vars[e]) > 1 else 0.0 for e in eps]

                t_means = [np.mean(epoch_text_vars[e]) if epoch_text_vars[e] else np.nan for e in eps]
                t_stds = [np.std(epoch_text_vars[e]) if len(epoch_text_vars[e]) > 1 else 0.0 for e in eps]

                m_means = [np.mean(epoch_mixed_vars[e]) if epoch_mixed_vars[e] else np.nan for e in eps]
                m_stds = [np.std(epoch_mixed_vars[e]) if len(epoch_mixed_vars[e]) > 1 else 0.0 for e in eps]

                eps, a_means, a_stds = np.array(eps), np.array(a_means), np.array(a_stds)
                t_means, t_stds, m_means, m_stds = np.array(t_means), np.array(t_stds), np.array(m_means), np.array(m_stds)

                # Audio Variance
                valid_a = ~np.isnan(a_means)
                if np.any(valid_a):
                    plt.plot(eps[valid_a], a_means[valid_a], label="Audio Representation Variance", color="blue", marker="o", linewidth=2)
                    plt.fill_between(eps[valid_a], a_means[valid_a] - a_stds[valid_a], a_means[valid_a] + a_stds[valid_a], color="blue", alpha=0.2)

                # Text Variance
                valid_t = ~np.isnan(t_means)
                if np.any(valid_t):
                    plt.plot(eps[valid_t], t_means[valid_t], label="Text Representation Variance", color="red", marker="s", linewidth=2)
                    plt.fill_between(eps[valid_t], t_means[valid_t] - t_stds[valid_t], t_means[valid_t] + t_stds[valid_t], color="red", alpha=0.2)

                # Mixed Variance
                valid_m = ~np.isnan(m_means)
                if np.any(valid_m):
                    plt.plot(eps[valid_m], m_means[valid_m], label="Mixed Representation Variance", color="purple", marker="^", linewidth=2)
                    plt.fill_between(eps[valid_m], m_means[valid_m] - m_stds[valid_m], m_means[valid_m] + m_stds[valid_m], color="purple", alpha=0.2)

            plt.xlabel("Epoch")
            plt.ylabel("State Variance (Trace of Covariance)")
            plt.title("Representation Variance Trajectory with Confidence Bounds (±1 std)")
            plt.grid(True, linestyle="--", alpha=0.6)
            plt.legend(loc="best")
            plt.tight_layout()
            plt.savefig(self.out_var_plot.get_path(), dpi=300, bbox_inches="tight")
            plt.close()
        except Exception:
            pass
